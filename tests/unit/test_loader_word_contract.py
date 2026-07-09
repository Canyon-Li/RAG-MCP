"""Unit tests for Word (DOCX) Loader contract and behavior.

Mirrors tests/unit/test_loader_pdf_contract.py:
- WordLoader initialization and configuration
- Input validation (extension / missing files)
- BaseLoader shared helpers (inherited in J2)
- Core DOCX conversion using fixtures generated on-the-fly with python-docx

Fixtures are generated dynamically into tmp_path (self-contained, no checked-in
binaries). If python-docx is unavailable the conversion tests are skipped; the
contract / validation / helper tests still run (they don't need real DOCX files).
"""

from pathlib import Path

import pytest

from src.core.types import Document
from src.libs.loader.base_loader import BaseLoader
from src.libs.loader.word_loader import WordLoader


# ---------------------------------------------------------------------------
# Fixtures: generate real .docx files with python-docx (skipped if absent)
# ---------------------------------------------------------------------------
@pytest.fixture
def simple_docx(tmp_path):
    pytest.importorskip("docx")
    import docx

    path = tmp_path / "simple.docx"
    doc = docx.Document()
    doc.add_heading("Simple Document", level=1)
    doc.add_paragraph("This is a test paragraph for the WordLoader.")
    doc.save(str(path))
    return path


@pytest.fixture
def images_docx(tmp_path):
    pytest.importorskip("docx")
    import docx
    from PIL import Image as PILImage

    path = tmp_path / "with_images.docx"
    src_img = tmp_path / "src.png"
    PILImage.new("RGB", (16, 16), (255, 0, 0)).save(str(src_img))

    doc = docx.Document()
    doc.add_heading("Document with Images", level=1)
    doc.add_paragraph("Below is an embedded image.")
    doc.add_picture(str(src_img))
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Initialization
# ---------------------------------------------------------------------------
class TestWordLoaderInitialization:
    """Tests for WordLoader initialization."""

    def test_default_initialization(self):
        """WordLoader can be initialized with defaults."""
        loader = WordLoader()
        assert loader.extract_images is True
        assert loader.image_storage_dir == Path("data/images")

    def test_custom_initialization(self, tmp_path):
        """WordLoader respects custom configuration."""
        loader = WordLoader(
            extract_images=False,
            image_storage_dir=str(tmp_path / "imgs"),
        )
        assert loader.extract_images is False
        assert loader.image_storage_dir == tmp_path / "imgs"

    def test_markitdown_available(self):
        """WordLoader requires MarkItDown to be available."""
        loader = WordLoader()
        assert loader._markitdown is not None


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
class TestWordLoaderValidation:
    """Tests for input validation."""

    def test_load_requires_docx_extension_txt(self, tmp_path):
        """load() raises ValueError for non-DOCX files (.txt)."""
        txt = tmp_path / "test.txt"
        txt.write_text("not a docx")
        with pytest.raises(ValueError, match="not a DOCX"):
            WordLoader().load(txt)

    def test_load_requires_docx_extension_pdf(self, tmp_path):
        """load() raises ValueError for .pdf files."""
        pdf = tmp_path / "test.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")
        with pytest.raises(ValueError, match="not a DOCX"):
            WordLoader().load(pdf)

    def test_load_nonexistent_file(self):
        """load() raises FileNotFoundError for missing files."""
        with pytest.raises(FileNotFoundError):
            WordLoader().load("nonexistent.docx")


# ---------------------------------------------------------------------------
# Inherited BaseLoader helpers (J2: moved to BaseLoader)
# ---------------------------------------------------------------------------
class TestWordLoaderHelperMethods:
    """Helpers inherited from BaseLoader are available on WordLoader."""

    def test_compute_file_hash_consistency(self, tmp_path):
        """_compute_file_hash returns consistent hash for same content."""
        f = tmp_path / "a.docx"
        f.write_bytes(b"consistent content")
        loader = WordLoader()
        assert loader._compute_file_hash(f) == loader._compute_file_hash(f)
        assert len(loader._compute_file_hash(f)) == 64  # SHA256 hex length

    def test_compute_file_hash_differs_for_different_content(self, tmp_path):
        """_compute_file_hash returns different hashes for different content."""
        f1, f2 = tmp_path / "1.docx", tmp_path / "2.docx"
        f1.write_bytes(b"content 1")
        f2.write_bytes(b"content 2")
        loader = WordLoader()
        assert loader._compute_file_hash(f1) != loader._compute_file_hash(f2)

    def test_extract_title_from_markdown_heading(self):
        """_extract_title finds first Markdown heading."""
        loader = WordLoader()
        assert loader._extract_title("# Main Title\n\nbody") == "Main Title"

    def test_extract_title_from_first_line(self):
        """_extract_title uses first non-empty line as fallback."""
        loader = WordLoader()
        assert loader._extract_title("First Line\n\nbody") == "First Line"

    def test_extract_title_handles_empty_text(self):
        """_extract_title returns None for empty text."""
        loader = WordLoader()
        assert loader._extract_title("") is None

    def test_generate_image_id_format(self):
        """_generate_image_id creates the expected id format (inherited static)."""
        # staticmethod inherited from BaseLoader is callable on the subclass
        assert WordLoader._generate_image_id("abc123def456", 2, 0) == "abc123de_2_0"


# ---------------------------------------------------------------------------
# Core DOCX conversion (real fixtures)
# ---------------------------------------------------------------------------
class TestDocxConversionCore:
    """Tests for core DOCX conversion functionality using real .docx files."""

    def test_convert_simple_docx_to_text(self, simple_docx, tmp_path):
        """Convert simple DOCX to text - verifies core Markdown conversion."""
        loader = WordLoader(
            extract_images=True,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = loader.load(simple_docx)

        assert isinstance(doc, Document)
        assert doc.id.startswith("doc_")
        assert len(doc.text) > 0
        assert doc.metadata["source_path"] == str(simple_docx)
        assert doc.metadata["doc_type"] == "docx"
        assert "doc_hash" in doc.metadata

    def test_extract_title_from_docx(self, simple_docx, tmp_path):
        """Verify title extraction from real DOCX (H1 heading)."""
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        doc = loader.load(simple_docx)
        assert doc.metadata.get("title") == "Simple Document"

    def test_docx_with_images_structure(self, images_docx, tmp_path):
        """Verify DOCX with images is processed with image extraction."""
        loader = WordLoader(
            extract_images=True,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = loader.load(images_docx)

        assert isinstance(doc, Document)
        assert len(doc.text) > 0
        assert "images" in doc.metadata
        images = doc.metadata["images"]
        assert isinstance(images, list)
        assert len(images) > 0

        for img in images:
            assert "id" in img
            assert "path" in img
            assert "page" in img
            assert "text_offset" in img
            assert "text_length" in img
            assert "position" in img
            assert Path(img["path"]).exists(), f"image file missing: {img['path']}"
            placeholder = f"[IMAGE: {img['id']}]"
            assert placeholder in doc.text, f"placeholder {placeholder} not in text"

    def test_image_extraction_disabled(self, images_docx, tmp_path):
        """Verify image extraction can be disabled."""
        loader = WordLoader(
            extract_images=False,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = loader.load(images_docx)
        assert len(doc.text) > 0
        assert "images" not in doc.metadata or doc.metadata.get("images") == []

    def test_document_hash_consistency(self, simple_docx, tmp_path):
        """Same DOCX produces same document hash (idempotency)."""
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        d1 = loader.load(simple_docx)
        d2 = loader.load(simple_docx)
        assert d1.metadata["doc_hash"] == d2.metadata["doc_hash"]
        assert d1.id == d2.id

    def test_document_serialization(self, simple_docx, tmp_path):
        """Loaded document can be serialized and recreated."""
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        doc = loader.load(simple_docx)

        d = doc.to_dict()
        assert isinstance(d, dict)
        assert d["metadata"]["doc_type"] == "docx"

        recreated = Document.from_dict(d)
        assert recreated.id == doc.id
        assert recreated.text == doc.text
