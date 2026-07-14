"""Unit tests for Word (DOCX) Parser contract and behavior.

Mirrors tests/unit/test_parser_pdf_contract.py:
- WordParser initialization and configuration
- Input validation (extension / missing files)
- BaseParser shared helpers (inherited)
- Core DOCX conversion using fixtures generated on-the-fly with python-docx

K1 (DEV_SPEC phase K): renamed from test_loader_word_contract.py. WordLoader
was renamed to WordParser; load() → parse(). Constructor contract unified but
all params have defaults so existing call styles still work.

Fixtures are generated dynamically into tmp_path (self-contained, no checked-in
binaries). If python-docx is unavailable the conversion tests are skipped; the
contract / validation / helper tests still run (they don't need real DOCX files).
"""

from pathlib import Path

import pytest

from src.core.types import Document
from src.libs.parser.base_parser import BaseParser
from src.libs.parser.word_parser import WordParser


# ---------------------------------------------------------------------------
# Fixtures: generate real .docx files with python-docx (skipped if absent)
# ---------------------------------------------------------------------------
@pytest.fixture
def simple_docx(tmp_path):
    pytest.importorskip("docx")
    import docx

    path = tmp_path / "simple.docx"
    doc = docx.Document()
    doc.add_heading("Simple Document", level=1)
    doc.add_paragraph("This is a test paragraph for the WordParser.")
    doc.save(str(path))
    return path


@pytest.fixture
def images_docx(tmp_path):
    pytest.importorskip("docx")
    import docx
    from PIL import Image as PILImage

    path = tmp_path / "with_images.docx"
    src_img = tmp_path / "src.png"
    PILImage.new("RGB", (16, 16), (255, 0, 0)).save(str(src_img))

    doc = docx.Document()
    doc.add_heading("Document with Images", level=1)
    doc.add_paragraph("Below is an embedded image.")
    doc.add_picture(str(src_img))
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Initialization
# ---------------------------------------------------------------------------
class TestWordParserInitialization:
    """Tests for WordParser initialization."""

    def test_default_initialization(self):
        """WordParser can be initialized with defaults."""
        parser = WordParser()
        assert parser.extract_images is True
        assert parser.image_storage_dir == Path("data/images")

    def test_custom_initialization(self, tmp_path):
        """WordParser respects custom configuration."""
        parser = WordParser(
            extract_images=False,
            image_storage_dir=str(tmp_path / "imgs"),
        )
        assert parser.extract_images is False
        assert parser.image_storage_dir == tmp_path / "imgs"

    def test_markitdown_available(self):
        """WordParser requires MarkItDown to be available."""
        parser = WordParser()
        assert parser._markitdown is not None


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------
class TestWordParserValidation:
    """Tests for input validation."""

    def test_parse_requires_docx_extension_txt(self, tmp_path):
        """parse() raises ValueError for non-DOCX files (.txt)."""
        txt = tmp_path / "test.txt"
        txt.write_text("not a docx")
        with pytest.raises(ValueError, match="not a DOCX"):
            WordParser().parse(txt)

    def test_parse_requires_docx_extension_pdf(self, tmp_path):
        """parse() raises ValueError for .pdf files."""
        pdf = tmp_path / "test.pdf"
        pdf.write_bytes(b"%PDF-1.4 fake")
        with pytest.raises(ValueError, match="not a DOCX"):
            WordParser().parse(pdf)

    def test_parse_nonexistent_file(self):
        """parse() raises FileNotFoundError for missing files."""
        with pytest.raises(FileNotFoundError):
            WordParser().parse("nonexistent.docx")


# ---------------------------------------------------------------------------
# Inherited BaseParser helpers
# ---------------------------------------------------------------------------
class TestWordParserHelperMethods:
    """Helpers inherited from BaseParser are available on WordParser."""

    def test_compute_file_hash_consistency(self, tmp_path):
        """_compute_file_hash returns consistent hash for same content."""
        f = tmp_path / "a.docx"
        f.write_bytes(b"consistent content")
        parser = WordParser()
        assert parser._compute_file_hash(f) == parser._compute_file_hash(f)
        assert len(parser._compute_file_hash(f)) == 64  # SHA256 hex length

    def test_compute_file_hash_differs_for_different_content(self, tmp_path):
        """_compute_file_hash returns different hashes for different content."""
        f1, f2 = tmp_path / "1.docx", tmp_path / "2.docx"
        f1.write_bytes(b"content 1")
        f2.write_bytes(b"content 2")
        parser = WordParser()
        assert parser._compute_file_hash(f1) != parser._compute_file_hash(f2)

    def test_extract_title_from_markdown_heading(self):
        """_extract_title finds first Markdown heading."""
        parser = WordParser()
        assert parser._extract_title("# Main Title\n\nbody") == "Main Title"

    def test_extract_title_from_first_line(self):
        """_extract_title uses first non-empty line as fallback."""
        parser = WordParser()
        assert parser._extract_title("First Line\n\nbody") == "First Line"

    def test_extract_title_handles_empty_text(self):
        """_extract_title returns None for empty text."""
        parser = WordParser()
        assert parser._extract_title("") is None

    def test_generate_image_id_format(self):
        """_generate_image_id creates the expected id format (inherited static)."""
        # staticmethod inherited from BaseParser is callable on the subclass
        assert WordParser._generate_image_id("abc123def456", 2, 0) == "abc123de_2_0"


# ---------------------------------------------------------------------------
# Core DOCX conversion (real fixtures)
# ---------------------------------------------------------------------------
class TestDocxConversionCore:
    """Tests for core DOCX conversion functionality using real .docx files."""

    def test_convert_simple_docx_to_text(self, simple_docx, tmp_path):
        """Convert simple DOCX to text - verifies core Markdown conversion."""
        parser = WordParser(
            extract_images=True,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = parser.parse(simple_docx)

        assert isinstance(doc, Document)
        assert doc.id.startswith("doc_")
        assert len(doc.text) > 0
        assert doc.metadata["source_path"] == str(simple_docx)
        assert doc.metadata["doc_type"] == "docx"
        assert "doc_hash" in doc.metadata

    def test_extract_title_from_docx(self, simple_docx, tmp_path):
        """Verify title extraction from real DOCX (H1 heading)."""
        parser = WordParser(image_storage_dir=str(tmp_path / "images"))
        doc = parser.parse(simple_docx)
        assert doc.metadata.get("title") == "Simple Document"

    def test_docx_with_images_structure(self, images_docx, tmp_path):
        """Verify DOCX with images is processed with image extraction."""
        parser = WordParser(
            extract_images=True,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = parser.parse(images_docx)

        assert isinstance(doc, Document)
        assert len(doc.text) > 0
        assert "images" in doc.metadata
        images = doc.metadata["images"]
        assert isinstance(images, list)
        assert len(images) > 0

        for img in images:
            assert "id" in img
            assert "path" in img
            assert "page" in img
            assert "text_offset" in img
            assert "text_length" in img
            assert "position" in img
            assert Path(img["path"]).exists(), f"image file missing: {img['path']}"
            placeholder = f"[IMAGE: {img['id']}]"
            assert placeholder in doc.text, f"placeholder {placeholder} not in text"

    def test_image_extraction_disabled(self, images_docx, tmp_path):
        """Verify image extraction can be disabled."""
        parser = WordParser(
            extract_images=False,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = parser.parse(images_docx)
        assert len(doc.text) > 0
        assert "images" not in doc.metadata or doc.metadata.get("images") == []

    def test_document_hash_consistency(self, simple_docx, tmp_path):
        """Same DOCX produces same document hash (idempotency)."""
        parser = WordParser(image_storage_dir=str(tmp_path / "images"))
        d1 = parser.parse(simple_docx)
        d2 = parser.parse(simple_docx)
        assert d1.metadata["doc_hash"] == d2.metadata["doc_hash"]
        assert d1.id == d2.id

    def test_document_serialization(self, simple_docx, tmp_path):
        """Parsed document can be serialized and recreated."""
        parser = WordParser(image_storage_dir=str(tmp_path / "images"))
        doc = parser.parse(simple_docx)

        d = doc.to_dict()
        assert isinstance(d, dict)
        assert d["metadata"]["doc_type"] == "docx"

        recreated = Document.from_dict(d)
        assert recreated.id == doc.id
        assert recreated.text == doc.text
