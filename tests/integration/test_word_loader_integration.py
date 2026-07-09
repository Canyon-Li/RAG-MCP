"""Integration tests for WordLoader against real .docx files.

Also verifies LoaderFactory integration: the ``docx`` provider is registered
and ``LoaderFactory.create(..., provider=docx)`` returns a WordLoader, so the
pluggable wiring from J1 actually reaches the new loader.

Fixtures are generated on-the-fly with python-docx; the whole module is skipped
if python-docx is unavailable (WordLoader itself only needs markitdown + stdlib).
"""

from pathlib import Path
from unittest.mock import MagicMock

import pytest

# Needed to generate fixtures; skip the whole module if absent.
pytest.importorskip("docx")

import docx  # noqa: E402
from PIL import Image as PILImage  # noqa: E402

from src.core.types import Document  # noqa: E402
from src.libs.loader.loader_factory import LoaderFactory  # noqa: E402
from src.libs.loader.word_loader import WordLoader  # noqa: E402


@pytest.fixture
def simple_docx(tmp_path):
    path = tmp_path / "simple.docx"
    d = docx.Document()
    d.add_heading("Simple Document", level=1)
    d.add_paragraph("Integration test paragraph.")
    d.save(str(path))
    return path


@pytest.fixture
def images_docx(tmp_path):
    path = tmp_path / "with_images.docx"
    src = tmp_path / "src.png"
    PILImage.new("RGB", (32, 32), (0, 128, 255)).save(str(src))
    d = docx.Document()
    d.add_heading("Document with Images", level=1)
    d.add_picture(str(src))
    d.save(str(path))
    return path


class TestWordLoaderWithRealFiles:
    """End-to-end WordLoader behavior against real DOCX files."""

    def test_load_simple_docx(self, simple_docx, tmp_path):
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        doc = loader.load(simple_docx)

        assert isinstance(doc, Document)
        assert doc.metadata["doc_type"] == "docx"
        assert "doc_hash" in doc.metadata
        assert len(doc.text) > 0

    def test_load_docx_with_images(self, images_docx, tmp_path):
        loader = WordLoader(
            extract_images=True,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = loader.load(images_docx)

        images = doc.metadata.get("images", [])
        assert len(images) >= 1
        for img in images:
            assert Path(img["path"]).exists()
            assert f"[IMAGE: {img['id']}]" in doc.text

    def test_load_without_image_extraction(self, simple_docx, tmp_path):
        loader = WordLoader(
            extract_images=False,
            image_storage_dir=str(tmp_path / "images"),
        )
        doc = loader.load(simple_docx)
        assert "images" not in doc.metadata or doc.metadata.get("images") == []

    def test_document_is_serializable(self, simple_docx, tmp_path):
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        doc = loader.load(simple_docx)
        assert Document.from_dict(doc.to_dict()).id == doc.id

    def test_file_hash_consistency(self, simple_docx, tmp_path):
        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        assert loader.load(simple_docx).id == loader.load(simple_docx).id

    def test_different_files_different_hash(self, tmp_path):
        f1, f2 = tmp_path / "a.docx", tmp_path / "b.docx"
        d1 = docx.Document()
        d1.add_paragraph("AAA")
        d1.save(str(f1))
        d2 = docx.Document()
        d2.add_paragraph("BBB")
        d2.save(str(f2))

        loader = WordLoader(image_storage_dir=str(tmp_path / "images"))
        assert loader.load(f1).id != loader.load(f2).id

    def test_custom_image_storage_dir(self, images_docx, tmp_path):
        """Extracted images land under the configured image_storage_dir."""
        custom = tmp_path / "custom_images"
        loader = WordLoader(
            extract_images=True,
            image_storage_dir=str(custom),
        )
        doc = loader.load(images_docx)
        assert len(doc.metadata["images"]) >= 1

        for img in doc.metadata["images"]:
            # stored under custom/{doc_hash}/{image_id}.ext
            assert Path(img["path"]).resolve().is_relative_to(custom.resolve())


class TestLoaderFactoryDocxIntegration:
    """J2 wiring: docx provider registered and routable via LoaderFactory."""

    def test_docx_provider_registered(self):
        assert "docx" in LoaderFactory.list_providers()

    def test_factory_creates_word_loader(self):
        settings = MagicMock()
        settings.ingestion.loader.provider = "docx"
        settings.ingestion.loader.extract_images = True

        loader = LoaderFactory.create(settings, collection="test_coll")
        assert isinstance(loader, WordLoader)
        # image_storage_dir is derived from the collection argument
        assert "test_coll" in str(loader.image_storage_dir)

    def test_factory_docx_and_pdf_both_registered(self):
        providers = LoaderFactory.list_providers()
        assert "pdf" in providers
        assert "docx" in providers
