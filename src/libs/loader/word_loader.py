"""Word (DOCX) Loader implementation.

Mirrors PdfLoader's contract: MarkItDown for text → Markdown, plus embedded
image extraction with ``[IMAGE: {id}]`` placeholders (C1 image contract).
A .docx file is a ZIP archive whose images live under ``word/media/``; we read
them with the stdlib ``zipfile`` module (no python-docx dependency at runtime)
and append placeholders at the end of the text — the same simplification
PdfLoader uses for PDF page boundaries.

Graceful Degradation:
- If MarkItDown is unavailable → ImportError at construction.
- If image extraction fails → log a warning and continue text-only.

J2 (DEV_SPEC phase J). Registered as the ``docx`` provider in LoaderFactory.
"""

from __future__ import annotations

import io
import logging
import os
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

try:
    from markitdown import MarkItDown

    MARKITDOWN_AVAILABLE = True
except ImportError:
    MARKITDOWN_AVAILABLE = False

try:
    from PIL import Image

    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

# python-docx is used as a graceful fallback when MarkItDown's docx backend
# (mammoth) is not installed. Optional at runtime.
try:
    import docx as _docx_module  # noqa: F401  (presence flag only)

    DOCX_FALLBACK_AVAILABLE = True
except ImportError:
    DOCX_FALLBACK_AVAILABLE = False

from src.core.types import Document
from src.libs.loader.base_loader import BaseLoader

logger = logging.getLogger(__name__)


class WordLoader(BaseLoader):
    """DOCX loader using MarkItDown for text + zipfile for embedded images.

    This loader:
    1. Extracts text from .docx and converts to Markdown (MarkItDown).
    2. Extracts embedded images from ``word/media/`` to data/images/{doc_hash}/.
    3. Appends ``[IMAGE: {image_id}]`` placeholders in extraction order.
    4. Records image metadata in ``Document.metadata.images`` (C1 contract).

    Constructor contract (shared with PdfLoader so LoaderFactory can build both):
        extract_images: Enable/disable image extraction (default True).
        image_storage_dir: Base directory for image storage (default data/images).

    Graceful Degradation:
        If image extraction fails, logs a warning and continues text-only.
    """

    SUPPORTED_EXTENSIONS = {".docx"}
    DOC_TYPE = "docx"

    def __init__(
        self,
        extract_images: bool = True,
        image_storage_dir: str | Path = "data/images",
    ) -> None:
        """Initialize Word Loader.

        Args:
            extract_images: Whether to extract embedded images from the DOCX.
            image_storage_dir: Base directory for storing extracted images.

        Raises:
            ImportError: If MarkItDown is not installed.
        """
        if not MARKITDOWN_AVAILABLE:
            raise ImportError(
                "MarkItDown is required for WordLoader. "
                "Install with: pip install markitdown"
            )

        self.extract_images = extract_images
        self.image_storage_dir = Path(image_storage_dir)
        self._markitdown = MarkItDown()

    def load(self, file_path: str | Path) -> Document:
        """Load and parse a .docx file.

        Args:
            file_path: Path to the DOCX file.

        Returns:
            Document with Markdown text and metadata.

        Raises:
            FileNotFoundError: If the file doesn't exist.
            ValueError: If the file is not a .docx.
            RuntimeError: If text parsing fails critically.
        """
        path = self._validate_file(file_path)
        if path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"File is not a DOCX: {path}")

        doc_hash = self._compute_file_hash(path)
        doc_id = f"doc_{doc_hash[:16]}"

        # Parse text: prefer MarkItDown (consistent with PdfLoader); fall back
        # to python-docx when MarkItDown's docx backend (mammoth) is missing.
        text_content = self._extract_text(path)

        metadata: Dict[str, Any] = {
            "source_path": str(path),
            "doc_type": self.DOC_TYPE,
            "doc_hash": doc_hash,
        }

        title = self._extract_title(text_content)
        if title:
            metadata["title"] = title

        if self.extract_images:
            try:
                text_content, images_metadata = self._extract_and_process_images(
                    path, text_content, doc_hash
                )
                if images_metadata:
                    metadata["images"] = images_metadata
            except Exception as e:
                logger.warning(
                    f"Image extraction failed for {path}, continuing with text-only: {e}"
                )

        return Document(id=doc_id, text=text_content, metadata=metadata)

    def _extract_text(self, path: Path) -> str:
        """Extract Markdown text from the DOCX.

        Prefers MarkItDown (high-quality Markdown, consistent with PdfLoader).
        Falls back to python-docx when MarkItDown's docx backend is unavailable
        (e.g. mammoth not installed) — graceful degradation per project rules.
        """
        try:
            result = self._markitdown.convert(str(path))
            return result.text_content if hasattr(result, "text_content") else str(result)
        except Exception as e:
            if not DOCX_FALLBACK_AVAILABLE:
                logger.error(f"Failed to parse DOCX {path}: {e}")
                raise RuntimeError(f"DOCX parsing failed: {e}") from e
            logger.warning(
                f"MarkItDown could not parse {path} ({type(e).__name__}); "
                f"falling back to python-docx extraction"
            )
            return self._extract_text_with_docx(path)

    def _extract_text_with_docx(self, path: Path) -> str:
        """Best-effort Markdown extraction via python-docx (fallback path).

        Maps Heading 1..6 / Title styles to Markdown ``#`` headings; other
        paragraphs are emitted verbatim. Good enough for retrieval chunking.
        """
        import docx

        document = docx.Document(str(path))
        lines: List[str] = []
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = (para.style.name or "").lower() if para.style else ""
            if style_name == "title":
                lines.append(f"# {text}")
            elif style_name.startswith("heading"):
                level = 1
                parts = style_name.split()
                if len(parts) > 1 and parts[-1].isdigit():
                    level = max(1, min(6, int(parts[-1])))
                lines.append(f"{'#' * level} {text}")
            else:
                lines.append(text)
        return "\n\n".join(lines)

    def _extract_and_process_images(
        self,
        docx_path: Path,
        text_content: str,
        doc_hash: str,
    ) -> Tuple[str, List[Dict[str, Any]]]:
        """Extract embedded images from word/media/ and append placeholders.

        DOCX stores images under ``word/media/`` (image1.png, image2.jpeg, ...).
        Placeholders are appended at the end of the text in extraction order —
        the same simplification strategy PdfLoader uses for PDF page boundaries.

        Args:
            docx_path: Path to the DOCX file.
            text_content: Extracted Markdown text.
            doc_hash: Document hash, used as the image sub-directory name.

        Returns:
            Tuple of (modified_text, images_metadata_list). On any failure the
            original text is returned with an empty list (graceful degradation).
        """
        if not self.extract_images:
            return text_content, []

        images_metadata: List[Dict[str, Any]] = []
        modified_text = text_content

        try:
            image_dir = self.image_storage_dir / doc_hash
            image_dir.mkdir(parents=True, exist_ok=True)

            with zipfile.ZipFile(docx_path) as zf:
                media_names = sorted(
                    name for name in zf.namelist() if name.startswith("word/media/")
                )

                if not media_names:
                    logger.debug(f"No images found in {docx_path}")
                    return text_content, []

                for idx, media_name in enumerate(media_names):
                    try:
                        ext = os.path.splitext(media_name)[1].lstrip(".").lower() or "png"
                        image_bytes = zf.read(media_name)

                        image_id = self._generate_image_id(doc_hash, 1, idx + 1)
                        image_filename = f"{image_id}.{ext}"
                        image_path = image_dir / image_filename
                        image_path.write_bytes(image_bytes)

                        # Dimensions (best-effort)
                        width, height = 0, 0
                        if PIL_AVAILABLE:
                            try:
                                img = Image.open(io.BytesIO(image_bytes))
                                width, height = img.size
                            except Exception:
                                width, height = 0, 0

                        placeholder = f"[IMAGE: {image_id}]"
                        insert_position = len(modified_text)
                        modified_text += f"\n{placeholder}\n"

                        # Relative path when inside repo cwd, else absolute
                        try:
                            stored_path = image_path.relative_to(Path.cwd())
                        except ValueError:
                            stored_path = image_path.absolute()

                        images_metadata.append({
                            "id": image_id,
                            "path": str(stored_path),
                            "page": 1,  # DOCX has no page concept
                            "text_offset": insert_position + 1,
                            "text_length": len(placeholder),
                            "position": {
                                "width": width,
                                "height": height,
                                "page": 1,
                                "index": idx,
                            },
                        })
                        logger.debug(f"Extracted image {image_id} from {docx_path}")

                    except Exception as e:
                        logger.warning(f"Failed to extract image {media_name}: {e}")
                        continue

            if images_metadata:
                logger.info(f"Extracted {len(images_metadata)} images from {docx_path}")

            return modified_text, images_metadata

        except Exception as e:
            logger.warning(f"Image extraction failed for {docx_path}: {e}")
            # Graceful degradation: return original text without images
            return text_content, []
